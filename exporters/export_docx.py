import os
from io import BytesIO
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor

def build_docx(
    transcript_text: str,
    title: str = "Transkrip Wawancara",
    language_label: str = "Bahasa Indonesia",
    save_path: str = None
) -> bytes:
    """
    Membuat berkas Word (.docx) dengan format rapi dan profesional.
    Dapat menyimpan langsung ke disk jika save_path disediakan.
    """
    doc = Document()
    
    # Pengaturan margin halaman
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Judul Wawancara
    heading = doc.add_heading(title, level=1)
    for run in heading.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59)

    # Metadata Dokumen
    created_at = datetime.now().strftime("%d/%m/%Y %H:%M WIB")
    meta_p = doc.add_paragraph()
    meta_run = meta_p.add_run(f"Waktu Transkripsi: {created_at} | Bahasa: {language_label}")
    meta_run.font.name = "Calibri"
    meta_run.font.size = Pt(9.5)
    meta_run.font.italic = True
    meta_run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph("")  # Spasi pemisah

    # Konten Transkrip
    for line in transcript_text.splitlines():
        trimmed = line.strip()
        if not trimmed:
            doc.add_paragraph("")
            continue

        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(4)

        run = p.add_run(trimmed)
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(30, 41, 59)

    # Simpan ke disk jika diminta
    if save_path:
        os.makedirs(os.path.dirname(save_path), exist_ok=True)
        doc.save(save_path)

    # Kembalikan bytes untuk keperluan download di antarmuka web
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
